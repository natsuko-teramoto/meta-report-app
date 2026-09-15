from pathlib import Path
from pptx.util import Inches, Pt
from ppt.common.layout import add_blank_slide
from ppt.common.text import add_textbox, add_section_title
from ppt.common.theme import PAGE_LEFT, PAGE_TOP, SLIDE_TITLE_SIZE
from services.report_export_service import build_cumulative_gap_figure, build_monthly_dataframe, build_trend_figure, visible_metric_keys
from utils.chart_export import save_chart

def add_api_cumulative_gap_slide(prs, period_result, visibility, image_prefix):
    show_gap=visibility.get("sections",{}).get("reach_impression_gap",True)
    show_monthly=visibility.get("sections",{}).get("monthly_trend",True)
    metrics=visible_metric_keys(visibility)
    monthly_df=build_monthly_dataframe(period_result) if show_monthly and metrics else None
    gap_fig=build_cumulative_gap_figure((period_result or {}).get("cumulative",{}),visibility) if show_gap else None
    monthly_fig=None
    if monthly_df is not None and not monthly_df.empty:
        monthly_fig=build_trend_figure(monthly_df,metrics,x_column="month_label")
    if gap_fig is None and monthly_fig is None: return None
    slide=add_blank_slide(prs)
    add_textbox(slide,"配信開始からの累計・月次推移",PAGE_LEFT,PAGE_TOP,Inches(9),Inches(0.5),font_size=SLIDE_TITLE_SIZE,bold=True)
    if gap_fig is not None:
        add_section_title(slide,"累計 リーチ・インプレッション",PAGE_LEFT,Inches(0.88),Inches(4.2),Inches(0.3))
        p=save_chart(gap_fig,f"{image_prefix}_cumulative_reach_impressions",width=1200,height=360,scale=2)
        slide.shapes.add_picture(str(Path(p)),PAGE_LEFT,Inches(1.18),width=Inches(9.0),height=Inches(1.65))
        cumulative=(period_result or {}).get("cumulative",{}) or {}
        reach=float(cumulative.get("reach",0) or 0)
        impressions=float(cumulative.get("impressions",0) or 0)
        frequency=impressions/reach if reach > 0 else 0.0
        add_textbox(slide,"1人当たり平均表示回数",Inches(9.85),Inches(1.42),Inches(2.75),Inches(0.34),font_size=Pt(14),bold=True)
        add_textbox(slide,f"{frequency:.2f}回",Inches(9.85),Inches(1.78),Inches(2.75),Inches(0.62),font_size=Pt(26),bold=True)
    if monthly_fig is not None:
        add_section_title(slide,"配信開始からの月次推移",PAGE_LEFT,Inches(3.18),Inches(4.0),Inches(0.3))
        monthly_fig.update_layout(height=470,margin=dict(l=55,r=30,t=45,b=45),font=dict(size=15))
        monthly_fig.update_xaxes(tickfont=dict(size=14))
        monthly_fig.update_yaxes(tickfont=dict(size=14))
        p=save_chart(monthly_fig,f"{image_prefix}_monthly",width=1600,height=700,scale=2)
        slide.shapes.add_picture(str(Path(p)),PAGE_LEFT,Inches(3.50),width=Inches(12.0),height=Inches(3.55))
    return slide
