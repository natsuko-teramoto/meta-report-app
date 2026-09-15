from pathlib import Path

from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

PAGE_LEFT = Inches(0.55)
PAGE_TOP = Inches(0.35)

TITLE_SIZE = Pt(23)
SECTION_SIZE = Pt(14)
BODY_SIZE = Pt(11)
SMALL_SIZE = Pt(9)

TEXT_COLOR = RGBColor(38, 38, 38)
MUTED_COLOR = RGBColor(100, 100, 100)
BORDER_COLOR = RGBColor(220, 220, 220)
CARD_FILL = RGBColor(248, 249, 251)


def add_title(slide, text, top=PAGE_TOP):
    box = slide.shapes.add_textbox(
        PAGE_LEFT, top, Inches(12.0), Inches(0.45)
    )
    p = box.text_frame.paragraphs[0]
    p.text = str(text)
    p.font.size = TITLE_SIZE
    p.font.bold = True
    p.font.color.rgb = TEXT_COLOR
    return box


def add_section_label(slide, text, left, top, width):
    box = slide.shapes.add_textbox(
        left, top, width, Inches(0.3)
    )
    p = box.text_frame.paragraphs[0]
    p.text = str(text)
    p.font.size = SECTION_SIZE
    p.font.bold = True
    p.font.color.rgb = TEXT_COLOR
    return box


def add_info_card(
    slide,
    label,
    value,
    left,
    top,
    width,
    height=Inches(0.74),
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_FILL
    shape.line.color.rgb = BORDER_COLOR

    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.06)

    p1 = tf.paragraphs[0]
    p1.text = str(label)
    p1.font.size = SMALL_SIZE
    p1.font.color.rgb = MUTED_COLOR

    p2 = tf.add_paragraph()
    p2.text = str(value)
    p2.font.size = BODY_SIZE
    p2.font.bold = True
    p2.font.color.rgb = TEXT_COLOR

    return shape


def add_metric_card(
    slide,
    item,
    left,
    top,
    width,
    height=Inches(0.9),
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_FILL
    shape.line.color.rgb = BORDER_COLOR

    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.10)
    tf.margin_right = Inches(0.10)
    tf.margin_top = Inches(0.06)

    p1 = tf.paragraphs[0]
    p1.text = item["label"]
    p1.font.size = SMALL_SIZE
    p1.font.color.rgb = MUTED_COLOR

    p2 = tf.add_paragraph()
    p2.text = item["value"]
    p2.font.size = Pt(17)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_COLOR

    extras = [
        text
        for text in (
            item.get("extra"),
            item.get("change"),
        )
        if text
    ]
    if extras:
        p3 = tf.add_paragraph()
        p3.text = " / ".join(extras)
        p3.font.size = Pt(8.5)
        p3.font.color.rgb = MUTED_COLOR

    return shape


def add_picture(
    slide,
    image_path,
    left,
    top,
    width=None,
    height=None,
):
    if not image_path:
        return None
    path = Path(image_path)
    if not path.exists():
        return None

    kwargs = {}
    if width is not None:
        kwargs["width"] = width
    if height is not None:
        kwargs["height"] = height

    return slide.shapes.add_picture(
        str(path), left, top, **kwargs
    )


def add_dataframe_table(
    slide,
    df,
    left,
    top,
    width,
    height,
    max_rows=None,
):
    if df is None or df.empty:
        return None

    data = df.copy()
    if max_rows is not None:
        data = data.head(max_rows)

    rows = len(data) + 1
    cols = len(data.columns)

    table = slide.shapes.add_table(
        rows, cols, left, top, width, height
    ).table

    for col_index, column in enumerate(data.columns):
        cell = table.cell(0, col_index)
        cell.text = str(column)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(242, 244, 247)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.size = Pt(8.5)
            paragraph.alignment = PP_ALIGN.CENTER

    for row_index, (_, row) in enumerate(data.iterrows(), start=1):
        for col_index, value in enumerate(row.tolist()):
            cell = table.cell(row_index, col_index)
            cell.text = str(value)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(8)
                paragraph.alignment = PP_ALIGN.CENTER

    return table
