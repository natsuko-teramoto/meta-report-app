from pathlib import Path
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from ppt.common.layout import add_blank_slide
from ppt.common.text import add_textbox, add_section_title
from ppt.common.theme import PAGE_LEFT, PAGE_TOP, SLIDE_TITLE_SIZE, TABLE_HEADER_FILL, TABLE_HEADER_TEXT_COLOR, TABLE_ALT_ROW_FILL, TABLE_BODY_FILL, COLOR_TEXT, FONT_NAME
from services.report_export_service import build_placement_bar_figure, build_placement_dataframe, build_placement_table
from utils.chart_export import save_chart

def _visible_metric_exists(visibility):
    return any(visibility.get("metrics", {}).get(k, True) for k in ["impressions","reach","clicks","link_clicks","landing_page_views"])

def _add_readable_table(slide, df, left, top, width, height, max_rows=10):
    data=df.head(max_rows); rows=len(data)+1; cols=len(data.columns)
    table=slide.shapes.add_table(rows,cols,left,top,width,height).table
    weights={"配置":1.55,"インプレッション":1.35,"リーチ":1.05,"クリック(すべて)":1.45,"リンククリック":1.35,"LPビュー":1.15}
    ws=[weights.get(c,1.0) for c in data.columns]; total=sum(ws)
    for i,w in enumerate(ws): table.columns[i].width=int(width*w/total)
    for c,name in enumerate(data.columns):
        cell=table.cell(0,c); cell.text=str(name); cell.fill.solid(); cell.fill.fore_color.rgb=TABLE_HEADER_FILL; cell.vertical_anchor=MSO_VERTICAL_ANCHOR.MIDDLE
        p=cell.text_frame.paragraphs[0]; p.font.name=FONT_NAME; p.font.size=Pt(13); p.font.bold=True; p.font.color.rgb=TABLE_HEADER_TEXT_COLOR; p.alignment=PP_ALIGN.CENTER
    for r,(_,row) in enumerate(data.iterrows(),start=1):
        for c,name in enumerate(data.columns):
            cell=table.cell(r,c); cell.text=str(row[name]); cell.fill.solid(); cell.fill.fore_color.rgb=TABLE_ALT_ROW_FILL if r%2==0 else TABLE_BODY_FILL; cell.vertical_anchor=MSO_VERTICAL_ANCHOR.MIDDLE
            p=cell.text_frame.paragraphs[0]; p.font.name=FONT_NAME; p.font.size=Pt(12.5); p.font.color.rgb=COLOR_TEXT; p.alignment=PP_ALIGN.LEFT if name=="配置" else PP_ALIGN.RIGHT
    return table

def add_api_placement_slide(prs, period_result, visibility, image_prefix, period_label):
    pv=visibility.get("placement",{}); show_chart=pv.get("chart",True); show_detail=pv.get("detail",True) and _visible_metric_exists(visibility)
    if not show_chart and not show_detail: return None
    df=build_placement_dataframe((period_result or {}).get("target_placement",[]))
    if df.empty: return None
    slide=add_blank_slide(prs)
    add_textbox(slide,f"表示場所分析：{period_label}",PAGE_LEFT,PAGE_TOP,Inches(8),Inches(0.5),font_size=SLIDE_TITLE_SIZE,bold=True)
    if show_chart:
        specs=[("impressions","インプレッション"),("reach","リーチ"),("clicks","クリック(すべて)")]
        cw=Inches(3.9); gap=Inches(0.25); top=Inches(1.20)
        for i,(key,label) in enumerate(specs):
            fig=build_placement_bar_figure(df,key,label)
            if fig is None: continue
            path=save_chart(fig,f"{image_prefix}_placement_{key}",width=1100,height=650,scale=2)
            left=PAGE_LEFT+i*(cw+gap)
            add_section_title(slide,label,left,Inches(0.86),cw,Inches(0.3))
            slide.shapes.add_picture(str(Path(path)),left,top,width=cw,height=Inches(2.70))
        detail_title_top=Inches(4.25); detail_table_top=Inches(4.68)
    else:
        detail_title_top=Inches(1.10); detail_table_top=Inches(1.53)
    if show_detail:
        table_df=build_placement_table(df,visibility)
        if not table_df.empty:
            add_section_title(slide,"詳細",PAGE_LEFT,detail_title_top,Inches(3),Inches(0.3))
            _add_readable_table(slide,table_df,PAGE_LEFT,detail_table_top,Inches(12.45),Inches(2.20),max_rows=10)
    return slide
