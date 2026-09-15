from pptx.util import Inches

from ppt.api_pages.common import (
    add_dataframe_table,
    add_picture,
    add_section_label,
    add_title,
)
from services.report_export_service import (
    build_monthly_dataframe,
    build_monthly_table,
    build_trend_figure,
    visible_metric_keys,
)
from utils.chart_export import save_chart


def add_api_monthly_slide(
    prs,
    period_result,
    visibility,
    image_prefix,
):
    if not visibility.get(
        "sections", {}
    ).get("monthly_trend", True):
        return None

    metrics = visible_metric_keys(visibility)
    if not metrics:
        return None

    df = build_monthly_dataframe(period_result)
    if df.empty:
        return None

    fig = build_trend_figure(
        df,
        metrics,
        x_column="month_label",
    )
    if fig is None:
        return None

    path = save_chart(
        fig,
        f"{image_prefix}_monthly",
        width=1600,
        height=850,
        scale=2,
    )

    table_df = build_monthly_table(
        df,
        metrics,
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "配信からの月次推移")

    add_picture(
        slide,
        path,
        Inches(0.55),
        Inches(1.00),
        width=Inches(12.15),
        height=Inches(3.65),
    )

    if not table_df.empty:
        add_section_label(
            slide,
            "月次実績",
            Inches(0.55),
            Inches(4.82),
            Inches(3),
        )
        add_dataframe_table(
            slide,
            table_df,
            Inches(0.55),
            Inches(5.15),
            Inches(12.15),
            Inches(2.05),
            max_rows=5,
        )

    return slide
